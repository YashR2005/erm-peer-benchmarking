import json

import pytest
from openpyxl import load_workbook

from erm_bench.pipeline import Pipeline, ScopeError
from erm_bench.schemas import ExtractionDraft, Request

from .fakes import FakeLLM, fake_fetch


@pytest.fixture
def run(tmp_path, setup):
    llm = FakeLLM()
    p = Pipeline.start(Request(client_name="City Developments", reporting_year=2025, peer_count=2),
                       setup, llm_factory=lambda s, store: llm, runs_dir=tmp_path,
                       fetch=fake_fetch, max_workers=2)
    p.research()
    return p, llm


def test_research_builds_scope(run):
    p, _ = run
    assert p.store.status == "awaiting_scope"

    profile = p.profile
    assert profile.bucket == "large"            # S$6.5B against provisional buckets
    assert profile.sector_guide == "real_estate"
    assert [f.sourced for f in profile.facts] == [True, True, False]
    assert any("CDL Hospitality Trusts" in w for w in profile.warnings)

    peers = p.peers
    assert [c.name for c in peers.candidates] == ["Alpha Land", "Beta Properties", "Gamma REIT"]
    assert [c.selected for c in peers.candidates] == [True, True, False]   # client dropped
    assert [c.bucket for c in peers.candidates] == ["large", "mid", "mid"]

    docs = {d.company_id: d for d in p.documents}
    assert docs["client"].status == "ok" and docs["client"].risk_pages == [4, 5]
    assert docs["client"].checks.year_match and docs["client"].checks.company_match
    assert docs["P2"].checks.has_risk_section            # found by page heading, no bookmarks
    assert docs["P3"].status == "not_pdf"
    assert [p.availability(c) for c in ["client", "P1", "P2", "P3"]] == [
        "good", "good", "good", "none"]


def test_checkpoint_1_edits_and_lock(run):
    p, _ = run
    assert [c.id for c in p.select(["gamma"])] == ["P3"]
    p.select(["P3"], selected=False)
    with pytest.raises(ScopeError):
        p.mark_na("D-nope-1", "x")
    p.mark_na("D-P3-1", "REIT; no annual report PDF")

    scope = p.lock_scope()
    assert scope["peer_ids"] == ["P1", "P2"]
    assert scope["documents"] == {"client": "D-client-1", "P1": "D-P1-1", "P2": "D-P2-1"}
    assert p.store.status == "analysing"
    with pytest.raises(ScopeError):
        p.select(["P3"])                          # scope is locked

    actions = [r["action"] for r in p.store.read_log("reviews.jsonl")]
    assert actions == ["select", "deselect", "mark_na", "lock_scope"]


def test_analysis_matrix_gaps_and_structure(run):
    p, _ = run
    p.lock_scope()
    p.analyse()
    f = p.findings()
    assert p.store.status == "awaiting_findings"

    risks = {r.id: r for r in f.risks}
    assert not risks["R-client-4"].verification.verified          # invented quote
    assert risks["R-P1-2"].verification.verified and risks["R-P1-2"].page == 3   # page corrected
    assert risks["R-P2-2"].theme_key == "unmapped"

    rows = {r.theme.key: r for r in f.rows}
    assert set(rows) == {"interest_rate", "property_market", "cyber", "construction_delay",
                         "workplace_safety", "fx", "unmapped"}
    assert [r.id for r in rows["property_market"].cells["client"]] == ["R-client-2"]
    assert all(rows["interest_rate"].discloses(c) for c in ["client", "P1", "P2"])
    assert rows["fx"].context_risk_ids == ["C2"] and not rows["fx"].discloses("client")

    gaps = {(g.type, g.theme.key) for g in f.gaps}
    assert ("relevant_not_disclosed", "fx") in gaps
    assert ("disclosed_not_in_context", "cyber") in gaps
    assert ("peers_disclose_client_silent", "workplace_safety") in gaps
    assert ("relevant_not_disclosed", "interest_rate") not in gaps

    obs = " ".join(f.structure.observations)
    assert "CDL uses 4 risk categories" in obs and "Alpha Land uses 5 risk categories" in obs
    assert "No company has a category for Sustainability" in obs
    assert "Safety & Security" in obs                 # outside the standard taxonomy
    assert "Beta Properties: Risks are listed without categories" in obs

    ctx = {c.id: c for c in f.context_risks}
    assert ctx["C2"].driver_fact_ids == ["F2"] and ctx["C2"].warnings


def test_checkpoint_2_edits_and_export(run):
    p, _ = run
    p.lock_scope()
    p.analyse()
    p.remove_themes(["unmapped"])
    p.rename_theme("cyber", "Cyber and data")
    p.recategorise_theme("fx", "strategic")
    with pytest.raises(ScopeError):
        p.recategorise_theme("fx", "nonsense")
    p.reject_evidence(["r-p2-1"])
    p.add_remark("Client could add People and Sustainability categories.")

    f = p.findings()
    rows = {r.theme.key: r for r in f.rows}
    assert "unmapped" not in rows
    assert rows["cyber"].theme.name == "Cyber and data"
    assert rows["fx"].theme.category == "strategic"
    assert not rows["interest_rate"].discloses("P2")

    path = p.export()
    assert p.store.status == "exported"
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Peer selection", "Risk coverage", "Gaps",
                             "Category structure", "Evidence", "Context risks", "Client profile",
                             "Documents"]
    peer_sheet = [[c.value for c in row] for row in wb["Peer selection"].iter_rows(values_only=False)]
    assert peer_sheet[1][0] == "CDL (client)"
    assert [r[12] for r in peer_sheet[2:]] == ["Y", "Y", "N"]
    coverage = [[c.value for c in row] for row in wb["Risk coverage"].iter_rows()]
    cyber = next(r for r in coverage if r[1] == "Cyber and data")
    assert cyber[4].startswith("Disclosed: R-client-3 (p.5)") and cyber[5] == "Not publicly evident"
    evidence = [[c.value for c in row] for row in wb["Evidence"].iter_rows()]
    invented = next(r for r in evidence if r[0] == "R-client-4")
    assert invented[11] == "no"
    remarks = [r for r in wb["Category structure"].iter_rows(values_only=True)
               if r[0] == "Consultant remark"]
    assert remarks and "Sustainability" in remarks[0][1]

    # A later edit reopens the findings.
    p.restore_themes(["unmapped"])
    assert p.store.status == "awaiting_findings"
    log = [json.loads(line)["action"] for line in
           (p.store.dir / "reviews.jsonl").read_text().splitlines()]
    assert log[-1] == "restore_theme" and "export" in log


def test_client_without_disclosure_reads_insufficient_evidence(run):
    p, llm = run
    llm.extractions["client"] = ExtractionDraft(has_risk_disclosure=False, categories=[],
                                                category_structure_note="No risk section.",
                                                risks=[])
    p.lock_scope()
    p.analyse()
    f = p.findings()
    assert not f.client_disclosed
    assert f.gaps == []


def test_peer_without_document_is_insufficient_evidence(run):
    p, _ = run
    p.select(["P3"])
    p.lock_scope()
    p.analyse()
    ext = {e.company_id: e for e in p.findings().extractions}
    assert ext["P3"].doc_id is None and "insufficient evidence" in ext["P3"].note


def test_set_bucket_rescans_and_keeps_consultant_peers(run, setup):
    p, llm = run
    peers = p.peers
    peers.candidates[1].added_by = "consultant"       # pretend Beta was added by hand
    p.store.save("peers", peers)
    p.set_bucket("large")
    new = p.peers
    assert new.bucket_filter == "large"
    assert new.candidates[0].name == "Beta Properties" and new.candidates[0].id == "P2"
    assert [c.id for c in new.candidates[1:]] == ["P4", "P5"]   # Beta itself is deduplicated
    assert any("Outside the requested large-cap bucket" in w for w in new.candidates[2].warnings)
    assert "Only 1 candidate(s) found in the large-cap bucket" in new.scan_notes
    assert [c.selected for c in new.candidates] == [True, True, False]
    assert {d.company_id for d in p.documents} == {"client", "P2", "P4", "P5"}


def test_announcement_page_is_followed_to_its_pdf(run):
    p, _ = run
    doc = p.add_document("P3", url="https://sgx.test/ann/1")
    assert doc.status == "uploaded"
    assert doc.url == "https://sgx.test/ann/1/1_Gamma_REIT_Annual_Report_2025.pdf"
    assert doc.checks.company_match and doc.checks.has_risk_section
    assert p.availability("P3") == "good"
