"""Phase 1 output: the Product A Excel workbook (v1.2 step A7)."""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .product_a import CLIENT
from .store import now_iso

if TYPE_CHECKING:
    from .pipeline import Pipeline

HEADER_FILL = PatternFill("solid", fgColor="002C77")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

LIMITATIONS = [
    "This workbook compares what companies disclose publicly. It does not assess how well any "
    "company manages risk.",
    "A risk that is not publicly evident may still be managed. Absence of disclosure is not "
    "absence of risk management.",
    "No two organisations are the same; read peer comparisons with their differences in mind.",
    "Every disclosed risk links to a verbatim quote and page. Quotes that could not be verified "
    "against the PDF text are listed on the Evidence sheet but excluded from the matrix.",
]


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "B2"
    return ws


def write_workbook(p: Pipeline) -> Path:
    f = p.findings()
    profile, peers, setup = p.profile, p.peers, p.setup
    meta, scope = p.store.meta, p.scope
    names = {c.id: c.name for c in f.companies}
    docs_by_id = {d.id: d for d in p.documents}
    wb = Workbook()
    wb.remove(wb.active)

    # Summary
    ws = _sheet(wb, "Summary", ["Item", "Value"], [
        ["Client", f"{profile.legal_name} ({profile.listing_label()})"],
        ["Peers", ", ".join(c.name for c in f.companies if c.id != CLIENT)],
        ["Reporting year", p.wanted_year or "latest available"],
        ["Market-cap filter", peers.bucket_filter or "none"],
        ["Method", f"{meta['method']} ({meta['setup_version']})"],
        ["Model", f"{meta['model']} via {meta.get('provider', 'anthropic')}"],
        ["Run", p.store.run_id],
        ["Scope locked", scope["locked_at"]],
        ["Exported", now_iso()],
        ["Client has its own risk disclosure", "yes" if f.client_disclosed else
         "no: client column reads 'insufficient evidence', no client gaps reported"],
        *[["Limitation", text] for text in LIMITATIONS],
    ], [32, 110])

    # Peer selection, as in Marsh's Appendix A
    def peer_row(name, hq, segments, countries, markets, assets, revenue, ebitda, employees,
                 listing, cap, bucket, selected, rationale, limitation, added_by, sources):
        return [name, hq, segments, countries, markets, assets, revenue, ebitda, employees,
                listing, cap, bucket, selected, rationale, limitation, added_by, sources]

    rows = [peer_row(f"{profile.short_name} (client)", profile.headquarters,
                     ", ".join(profile.segments), len(profile.countries) or "",
                     ", ".join(profile.main_markets), profile.total_assets, profile.revenue,
                     profile.ebitda, profile.employees, profile.listing_label(),
                     profile.market_cap_sgd_millions, profile.bucket, "-", "-", "-", "-",
                     "\n".join(sorted({s.url for fct in profile.facts for s in fct.sources})))]
    for c in peers.candidates:
        rows.append(peer_row(
            c.name, c.headquarters, ", ".join(c.segments), c.operating_countries,
            ", ".join(c.main_markets), c.total_assets, c.revenue, c.ebitda, c.employees,
            c.listing_label(), c.market_cap_sgd_millions, c.bucket,
            "Y" if c.id in scope["peer_ids"] else "N", c.rationale, c.key_limitation,
            c.added_by, "\n".join(s.url for s in c.sources)))
    _sheet(wb, "Peer selection", [
        "Peer", "HQ", "Key operating segments", "# of operating countries", "Predominant markets",
        "Total assets", "Total revenue", "EBITDA", "# of employees", "Listing",
        f"Market cap ({setup.bucket_currency}M)", "Market-cap bucket", "Selected for analysis?",
        "Rationale", "Key limitation", "Added by", "Sources",
    ], rows, [26, 14, 30, 12, 20, 14, 14, 14, 12, 14, 12, 10, 10, 45, 35, 10, 50])

    # Risk coverage
    headers = ["Category", "Theme", "Description", "Context-inferred (client)"]
    headers += [names[c.id] for c in f.companies]
    rows = []
    for row in f.rows:
        record = [setup.taxonomy_name(row.theme.category), row.theme.name,
                  row.theme.description, ", ".join(row.context_risk_ids)]
        for c in f.companies:
            if c.id == CLIENT and not f.client_disclosed:
                record.append("Insufficient evidence")
                continue
            risks = row.cells.get(c.id) or []
            record.append("Disclosed: " + "; ".join(f"{r.id} (p.{r.page})" for r in risks)
                          if risks else "Not publicly evident")
        rows.append(record)
    _sheet(wb, "Risk coverage", headers, rows, [16, 32, 40, 16] + [24] * len(f.companies))

    # Gaps
    _sheet(wb, "Gaps", ["Gap type", "Theme", "Category", "Detail"],
           [[g.label, g.theme.name, setup.taxonomy_name(g.theme.category), g.detail]
            for g in f.gaps] or [["No gaps reported", "", "", "The client has no verified "
                                  "disclosure of its own" if not f.client_disclosed else ""]],
           [40, 32, 18, 40])

    # Category structure (Client X appendix p.15)
    rows = [["Number of categories"] + [
        len(col.categories) if col.has_structure else "none found" for col in f.structure.columns]]
    for key in [*setup.taxonomy_keys, "other"]:
        cells = [", ".join(f.structure.by_taxonomy[key][c.id]) or "—" for c in f.companies]
        if any(cell != "—" for cell in cells) or key != "other":
            rows.append([setup.taxonomy_name(key)] + cells)
    rows.append([""] * (1 + len(f.companies)))
    rows += [["Observation", o] + [""] * (len(f.companies) - 1) for o in f.structure.observations]
    rows += [["Consultant remark", r] + [""] * (len(f.companies) - 1) for r in f.remarks]
    ws = _sheet(wb, "Category structure", ["Standard category"] + [names[c.id] for c in f.companies],
                rows, [26] + [36] * len(f.companies))
    ws.append([])
    ws.append(["Company", "Category (own wording)", "Sub-categories", "Mapped to", "Quote", "Page",
               "Verified"])
    for e in f.extractions:
        for k in e.categories:
            ws.append([names.get(e.company_id, e.company_id), k.name, ", ".join(k.sub_categories),
                       setup.taxonomy_name(k.taxonomy_key or "other"), k.quote, k.page,
                       "yes" if k.verification.verified else "no"])

    # Evidence
    rows = []
    themes = {t.key: t for t in f.themes}
    for e in f.extractions:
        doc = docs_by_id.get(e.doc_id or "")
        for r in e.risks:
            theme = themes.get(r.theme_key or "")
            rows.append([r.id, names.get(r.company_id, r.company_id),
                         theme.name if theme else "", setup.taxonomy_name(theme.category)
                         if theme else "", r.title, r.summary, r.company_category, r.quote, r.page,
                         doc.title if doc else "", doc.url if doc else "",
                         "yes" if r.verification.verified else "no",
                         round(r.verification.score), r.verification.note,
                         "yes" if r.rejected else ""])
    _sheet(wb, "Evidence", [
        "Risk id", "Company", "Theme", "Category", "Company wording", "Summary",
        "Company category", "Verbatim quote", "Page", "Document", "URL", "Verified",
        "Match score", "Verification note", "Rejected by reviewer",
    ], rows, [12, 20, 28, 16, 30, 40, 20, 60, 6, 30, 40, 9, 9, 30, 10])

    # Context-inferred risks
    facts = {fct.id: fct.text for fct in profile.facts}
    _sheet(wb, "Context risks", [
        "Id", "Risk", "Category", "Driver", "Profile facts", "Consequence", "Affected", "Theme",
        "Warnings",
    ], [[c.id, c.title, setup.taxonomy_name(c.category), c.driver,
         "\n".join(f"{fid}: {facts.get(fid, '')}" for fid in c.driver_fact_ids), c.consequence,
         c.affected, themes[c.theme_key].name if c.theme_key in themes else "",
         "; ".join(c.warnings)] for c in f.context_risks],
        [6, 32, 16, 40, 50, 40, 20, 28, 30])

    # Client profile facts
    _sheet(wb, "Client profile", ["Id", "Field", "Fact", "Sources", "Cited text"], [
        [fct.id, fct.field, fct.text, "\n".join(s.url for s in fct.sources) or "unsourced",
         "\n".join(s.cited_text for s in fct.sources if s.cited_text)] for fct in profile.facts
    ], [6, 14, 60, 50, 60])

    # Documents
    ids = {CLIENT, *scope["peer_ids"]}
    _sheet(wb, "Documents", [
        "Doc", "Company", "Type", "Title", "FY", "Status", "Used", "Company check", "Year found",
        "Risk section", "Risk pages", "Note", "URL", "SHA-256",
    ], [[d.id, p.company_name(d.company_id), d.doc_type, d.title, d.fiscal_year, d.status,
         "yes" if scope["documents"].get(d.company_id) == d.id else "",
         d.checks.company_match, d.checks.year_found, d.checks.has_risk_section,
         ", ".join(map(str, d.risk_pages)), d.note, d.url, d.sha256]
        for d in p.documents if d.company_id in ids],
        [12, 20, 16, 30, 6, 12, 6, 10, 8, 8, 20, 40, 40, 20])

    path = p.store.exports_dir / f"{p.store.run_id}-product-a.xlsx"
    wb.save(path)
    return path
